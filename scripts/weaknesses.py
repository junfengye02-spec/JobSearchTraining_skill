#!/usr/bin/env python3
"""Create and maintain the single interview weakness review workbook."""

from __future__ import annotations

import argparse
import json
import os
import re
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as exc:  # pragma: no cover - environment dependent
    raise SystemExit("openpyxl is required. Install it with: python -m pip install openpyxl") from exc


RUNTIME_DIR = ".adaptive-interview-coach"
DETAIL_SHEET = "薄弱点明细"
REVIEW_SHEET = "复习记录"
DASHBOARD_SHEET = "复习看板"
OPTIONS_SHEET = "_选项"

DETAIL_HEADERS = [
    "编号", "首次发现日期", "最近更新日期", "分类", "知识点", "来源问题/岗位",
    "原回答薄弱表现", "正确结论", "原理讲解", "简历项目/代码示例", "大厂回答框架",
    "关键词", "常见追问与陷阱", "确认题", "我的复述/答案", "掌握状态",
    "掌握度(0-5)", "复习次数", "下次复习日期", "所属期/Day", "任务标题", "备注",
]
REVIEW_HEADERS = [
    "记录编号", "薄弱点编号", "复习日期", "复习方式", "复习前掌握度", "复习后掌握度",
    "验证题", "回答摘要", "是否通过", "下次复习日期", "所属期/Day", "任务标题", "备注",
]
CATEGORIES = [
    "项目深挖", "Java基础", "JVM", "并发编程", "Spring", "MySQL", "Redis", "RocketMQ",
    "微服务", "网络", "操作系统", "系统设计", "算法", "AI Agent", "RAG",
    "Function Calling", "岗位JD", "其他",
]
STATUSES = ["待讲解", "讲解中", "待验证", "需复习", "已掌握"]

NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT_BLUE = "EAF3F8"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D9E1F2")
FONT_NAME = "Microsoft YaHei"


def now_iso() -> str:
    return datetime.now().astimezone().isoformat()


def today_iso() -> str:
    return date.today().isoformat()


def normalize(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "").strip().lower())


def read_json(path: Path, default: object) -> object:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return default


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temp.replace(path)


def load_config(workspace: Path) -> dict:
    path = workspace / RUNTIME_DIR / "config.json"
    if not path.exists():
        raise SystemExit(f"Workspace is not initialized: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_path(workspace: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else workspace / path


def workbook_path(workspace: Path, config: dict | None = None) -> Path:
    config = config or load_config(workspace)
    return resolve_path(workspace, config["files"]["weakness_workbook"])


def pending_path(workspace: Path, config: dict | None = None) -> Path:
    config = config or load_config(workspace)
    return resolve_path(workspace, config["files"]["pending_weakness_updates"])


def style_header(sheet, headers: list[str]) -> None:
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(1, column, header)
        cell.font = Font(name=FONT_NAME, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(right=THIN)
    sheet.row_dimensions[1].height = 34


def create_dashboard(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = DASHBOARD_SHEET
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "面试薄弱点复习看板"
    sheet["A1"].font = Font(name=FONT_NAME, size=18, bold=True, color=WHITE)
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 34

    cards = [
        ("A3", "待处理", f'=COUNTIFS(\'{DETAIL_SHEET}\'!$P:$P,"<>已掌握",\'{DETAIL_SHEET}\'!$A:$A,"<>")'),
        ("C3", "今日到期", f'=COUNTIFS(\'{DETAIL_SHEET}\'!$S:$S,"<="&TODAY(),\'{DETAIL_SHEET}\'!$S:$S,">0",\'{DETAIL_SHEET}\'!$P:$P,"<>已掌握")'),
        ("E3", "已掌握", f'=COUNTIF(\'{DETAIL_SHEET}\'!$P:$P,"已掌握")'),
        ("A6", "总薄弱点", f'=MAX(COUNTA(\'{DETAIL_SHEET}\'!$A:$A)-1,0)'),
        ("C6", "平均掌握度", f'=IFERROR(AVERAGE(\'{DETAIL_SHEET}\'!$Q:$Q),0)'),
        ("E6", "累计复习次数", f'=SUM(\'{DETAIL_SHEET}\'!$R:$R)'),
    ]
    for address, label, formula in cards:
        cell = sheet[address]
        row, column = cell.row, cell.column
        sheet.merge_cells(start_row=row, start_column=column, end_row=row, end_column=column + 1)
        cell.value = label
        cell.font = Font(name=FONT_NAME, bold=True, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center")
        sheet.merge_cells(start_row=row + 1, start_column=column, end_row=row + 1, end_column=column + 1)
        value = sheet.cell(row + 1, column, formula)
        value.font = Font(name=FONT_NAME, size=18, bold=True, color=NAVY)
        value.alignment = Alignment(horizontal="center")
        value.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        value.number_format = "0.0" if label == "平均掌握度" else "0"

    sheet["A10"], sheet["B10"] = "分类", "薄弱点数量"
    for cell in sheet[10][:2]:
        cell.font = Font(name=FONT_NAME, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center")
    for index, category in enumerate(CATEGORIES, 11):
        sheet.cell(index, 1, category)
        sheet.cell(index, 2, f'=COUNTIF(\'{DETAIL_SHEET}\'!$D:$D,A{index})')
        for column in (1, 2):
            sheet.cell(index, column).border = Border(bottom=THIN)
            sheet.cell(index, column).font = Font(name=FONT_NAME)

    sheet.merge_cells("D10:F10")
    sheet["D10"] = "复习规则"
    sheet["D10"].font = Font(name=FONT_NAME, bold=True, color=WHITE)
    sheet["D10"].fill = PatternFill("solid", fgColor=NAVY)
    notes = [
        "发现薄弱点后立即记录，再进行讲解与确认。",
        "每次确认题、复述或复习都追加历史记录。",
        "0-2分次日复习，3-4分三天后，5分七天后。",
        "未确认掌握前不进入下一道面试主问题。",
    ]
    for index, note in enumerate(notes, 11):
        sheet.merge_cells(start_row=index, start_column=4, end_row=index, end_column=6)
        sheet.cell(index, 4, note)
        sheet.cell(index, 4).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(index, 4).font = Font(name=FONT_NAME)
    for column, width in {"A": 18, "B": 16, "C": 4, "D": 20, "E": 20, "F": 20}.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A10"
    sheet.sheet_view.showGridLines = False


def create_detail_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(DETAIL_SHEET)
    style_header(sheet, DETAIL_HEADERS)
    widths = [12, 14, 14, 16, 22, 34, 38, 38, 48, 42, 42, 24, 36, 34, 34, 14, 14, 12, 14, 16, 30, 28]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[chr(64 + index) if index <= 26 else "A"].width = width
    sheet.freeze_panes = "F2"
    sheet.auto_filter.ref = "A1:V5000"
    sheet.sheet_view.showGridLines = False

    category_validation = DataValidation(type="list", formula1=f"='{OPTIONS_SHEET}'!$A$1:$A${len(CATEGORIES)}", allow_blank=True)
    status_validation = DataValidation(type="list", formula1=f"='{OPTIONS_SHEET}'!$B$1:$B${len(STATUSES)}", allow_blank=True)
    score_validation = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True)
    sheet.add_data_validation(category_validation)
    sheet.add_data_validation(status_validation)
    sheet.add_data_validation(score_validation)
    category_validation.add("D2:D5000")
    status_validation.add("P2:P5000")
    score_validation.add("Q2:Q5000")
    sheet.conditional_formatting.add("P2:P5000", FormulaRule(formula=['$P2="已掌握"'], fill=PatternFill("solid", fgColor=GREEN)))
    sheet.conditional_formatting.add("P2:P5000", FormulaRule(formula=['OR($P2="待讲解",$P2="讲解中")'], fill=PatternFill("solid", fgColor=RED)))
    sheet.conditional_formatting.add("P2:P5000", FormulaRule(formula=['OR($P2="待验证",$P2="需复习")'], fill=PatternFill("solid", fgColor=YELLOW)))
    sheet.conditional_formatting.add("Q2:Q5000", FormulaRule(formula=['AND($Q2<3,$Q2<>"")'], fill=PatternFill("solid", fgColor=RED)))
    sheet.conditional_formatting.add("S2:S5000", FormulaRule(formula=['AND($S2<=TODAY(),$S2<>"",$P2<>"已掌握")'], fill=PatternFill("solid", fgColor=YELLOW)))


def create_review_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet(REVIEW_SHEET)
    style_header(sheet, REVIEW_HEADERS)
    widths = [14, 14, 14, 18, 16, 16, 36, 42, 14, 14, 16, 30, 28]
    from openpyxl.utils import get_column_letter
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:M5000"
    sheet.sheet_view.showGridLines = False
    pass_validation = DataValidation(type="list", formula1='"是,否,待确认"', allow_blank=True)
    before_validation = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True)
    after_validation = DataValidation(type="whole", operator="between", formula1="0", formula2="5", allow_blank=True)
    sheet.add_data_validation(pass_validation)
    sheet.add_data_validation(before_validation)
    sheet.add_data_validation(after_validation)
    pass_validation.add("I2:I5000")
    before_validation.add("E2:E5000")
    after_validation.add("F2:F5000")


def build_workbook() -> Workbook:
    workbook = Workbook()
    create_dashboard(workbook)
    create_detail_sheet(workbook)
    create_review_sheet(workbook)
    options = workbook.create_sheet(OPTIONS_SHEET)
    for index, category in enumerate(CATEGORIES, 1):
        options.cell(index, 1, category)
    for index, status in enumerate(STATUSES, 1):
        options.cell(index, 2, status)
    options.sheet_state = "hidden"
    return workbook


def safe_save(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name + ".tmp.xlsx")
    try:
        workbook.save(temp)
        os.replace(temp, path)
    except PermissionError as exc:
        temp.unlink(missing_ok=True)
        raise PermissionError(f"Workbook is open or locked: {path}") from exc
    finally:
        temp.unlink(missing_ok=True)


def ensure_workbook(workspace: Path) -> Path:
    config = load_config(workspace)
    path = workbook_path(workspace, config)
    if not path.exists():
        safe_save(build_workbook(), path)
    load_workbook(path, read_only=True).close()
    return path


def style_data_row(sheet, row: int, columns: int) -> None:
    for column in range(1, columns + 1):
        cell = sheet.cell(row, column)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.font = Font(name=FONT_NAME, size=10)
        cell.border = Border(bottom=THIN)
    for column in (2, 3, 19):
        if column <= columns:
            sheet.cell(row, column).number_format = "yyyy-mm-dd"
    sheet.row_dimensions[row].height = 78


def first_empty_row(sheet) -> int:
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row, 1).value in (None, ""):
            return row
    return max(2, sheet.max_row + 1)


def next_identifier(sheet, prefix: str) -> str:
    highest = 0
    pattern = re.compile(rf"^{re.escape(prefix)}-(\d+)$")
    for row in range(2, sheet.max_row + 1):
        match = pattern.match(str(sheet.cell(row, 1).value or ""))
        if match:
            highest = max(highest, int(match.group(1)))
    return f"{prefix}-{highest + 1:04d}"


def find_detail_row(sheet, data: dict) -> int | None:
    supplied_id = str(data.get("weakness_id") or data.get("id") or "").strip()
    category = normalize(data.get("category"))
    topic = normalize(data.get("topic"))
    for row in range(2, sheet.max_row + 1):
        if supplied_id and str(sheet.cell(row, 1).value or "") == supplied_id:
            return row
        if category and topic:
            if normalize(sheet.cell(row, 4).value) == category and normalize(sheet.cell(row, 5).value) == topic:
                return row
    return None


def int_score(value: object, default: int = 0) -> int:
    try:
        return max(0, min(5, int(value)))
    except (TypeError, ValueError):
        return default


DETAIL_FIELD_COLUMNS = {
    "category": 4,
    "topic": 5,
    "source": 6,
    "weakness": 7,
    "correct_conclusion": 8,
    "explanation": 9,
    "example": 10,
    "answer_framework": 11,
    "keywords": 12,
    "traps": 13,
    "confirmation_question": 14,
    "user_answer": 15,
    "status": 16,
    "mastery": 17,
    "review_count": 18,
    "next_review": 19,
    "period_day": 20,
    "task_title": 21,
    "notes": 22,
}


def apply_upsert(workbook: Workbook, data: dict) -> str:
    sheet = workbook[DETAIL_SHEET]
    row = find_detail_row(sheet, data)
    is_new = row is None
    if is_new:
        row = first_empty_row(sheet)
        sheet.cell(row, 1, next_identifier(sheet, "WP"))
        sheet.cell(row, 2, data.get("first_seen") or today_iso())
        sheet.cell(row, 18, 0)
    sheet.cell(row, 3, data.get("updated_at") or today_iso())
    if not data.get("status"):
        data = {**data, "status": "讲解中" if is_new else sheet.cell(row, 16).value}
    if "mastery" in data:
        data = {**data, "mastery": int_score(data.get("mastery"))}
    for field, column in DETAIL_FIELD_COLUMNS.items():
        if field not in data or data[field] is None:
            continue
        value = data[field]
        if isinstance(value, list):
            value = "、".join(str(item) for item in value)
        sheet.cell(row, column, value)
    style_data_row(sheet, row, len(DETAIL_HEADERS))
    return str(sheet.cell(row, 1).value)


def next_review_date(mastery: int) -> str:
    days = 1 if mastery <= 2 else 3 if mastery <= 4 else 7
    return (date.today() + timedelta(days=days)).isoformat()


def passed_text(value: object) -> str:
    if isinstance(value, bool):
        return "是" if value else "否"
    text = str(value or "").strip()
    if text in ("是", "否", "待确认"):
        return text
    return "是" if text.lower() in ("true", "pass", "passed", "yes") else "否"


def apply_review(workbook: Workbook, data: dict) -> str:
    detail = workbook[DETAIL_SHEET]
    row = find_detail_row(detail, data)
    if row is None:
        raise ValueError("Review references an unknown weakness. Upsert the weakness first.")
    weakness_id = str(detail.cell(row, 1).value)
    before = int_score(detail.cell(row, 17).value)
    after = int_score(data.get("mastery_after"), before)
    passed = passed_text(data.get("passed"))
    next_date = str(data.get("next_review") or next_review_date(after))
    detail.cell(row, 3, data.get("review_date") or today_iso())
    detail.cell(row, 15, data.get("answer_summary") or data.get("user_answer") or "")
    detail.cell(row, 16, "已掌握" if passed == "是" else "需复习")
    detail.cell(row, 17, after)
    detail.cell(row, 18, int(detail.cell(row, 18).value or 0) + 1)
    detail.cell(row, 19, next_date)
    if data.get("period_day"):
        detail.cell(row, 20, data["period_day"])
    if data.get("task_title"):
        detail.cell(row, 21, data["task_title"])
    style_data_row(detail, row, len(DETAIL_HEADERS))

    review = workbook[REVIEW_SHEET]
    review_row = first_empty_row(review)
    review.cell(review_row, 1, next_identifier(review, "RV"))
    review.cell(review_row, 2, weakness_id)
    review.cell(review_row, 3, data.get("review_date") or today_iso())
    review.cell(review_row, 4, data.get("method") or "确认题/复述")
    review.cell(review_row, 5, before)
    review.cell(review_row, 6, after)
    review.cell(review_row, 7, data.get("question") or detail.cell(row, 14).value or "")
    review.cell(review_row, 8, data.get("answer_summary") or data.get("user_answer") or "")
    review.cell(review_row, 9, passed)
    review.cell(review_row, 10, next_date)
    review.cell(review_row, 11, data.get("period_day") or detail.cell(row, 20).value or "")
    review.cell(review_row, 12, data.get("task_title") or detail.cell(row, 21).value or "")
    review.cell(review_row, 13, data.get("notes") or "")
    style_data_row(review, review_row, len(REVIEW_HEADERS))
    return weakness_id


def load_operations(path: Path, kind: str) -> list[dict]:
    value = read_json(path, [])
    values = value if isinstance(value, list) else [value]
    return [{"kind": kind, "data": item} for item in values if isinstance(item, dict)]


def queue_operations(path: Path, operations: list[dict]) -> None:
    existing = read_json(path, [])
    if not isinstance(existing, list):
        existing = []
    existing.extend({"queued_at": now_iso(), **operation} for operation in operations)
    write_json(path, existing)


def apply_operations(workspace: Path, operations: list[dict]) -> dict:
    config = load_config(workspace)
    book_path = ensure_workbook(workspace)
    queue_path = pending_path(workspace, config)
    pending = read_json(queue_path, [])
    pending_ops = [
        {"kind": item.get("kind"), "data": item.get("data", {})}
        for item in pending
        if isinstance(item, dict)
    ] if isinstance(pending, list) else []
    all_operations = pending_ops + operations
    if not all_operations:
        return {"applied": [], "pending": 0}

    workbook = load_workbook(book_path)
    applied: list[dict] = []
    try:
        for operation in all_operations:
            kind = operation.get("kind")
            data = operation.get("data", {})
            identifier = apply_upsert(workbook, data) if kind == "upsert" else apply_review(workbook, data)
            applied.append({"kind": kind, "id": identifier})
        safe_save(workbook, book_path)
        load_workbook(book_path, read_only=True).close()
    except PermissionError:
        queue_operations(queue_path, operations)
        return {"applied": [], "pending": len(pending_ops) + len(operations), "locked": True}
    if queue_path.exists():
        queue_path.unlink()
    return {"applied": applied, "pending": 0, "workbook": str(book_path)}


def cmd_init(args: argparse.Namespace) -> int:
    path = ensure_workbook(Path(args.workspace).expanduser().resolve())
    print(json.dumps({"workbook": str(path)}, ensure_ascii=False, indent=2))
    return 0


def cmd_upsert(args: argparse.Namespace) -> int:
    workspace = Path(args.workspace).expanduser().resolve()
    result = apply_operations(workspace, load_operations(Path(args.input).expanduser().resolve(), "upsert"))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 3 if result.get("locked") else 0


def cmd_review(args: argparse.Namespace) -> int:
    workspace = Path(args.workspace).expanduser().resolve()
    result = apply_operations(workspace, load_operations(Path(args.input).expanduser().resolve(), "review"))
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 3 if result.get("locked") else 0


def cmd_flush(args: argparse.Namespace) -> int:
    result = apply_operations(Path(args.workspace).expanduser().resolve(), [])
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 3 if result.get("locked") else 0


def cmd_summary(args: argparse.Namespace) -> int:
    workspace = Path(args.workspace).expanduser().resolve()
    path = ensure_workbook(workspace)
    workbook = load_workbook(path, read_only=True, data_only=False)
    sheet = workbook[DETAIL_SHEET]
    counts = {status: 0 for status in STATUSES}
    total = 0
    due = 0
    for row in range(2, sheet.max_row + 1):
        if not sheet.cell(row, 1).value:
            continue
        total += 1
        status = str(sheet.cell(row, 16).value or "")
        counts[status] = counts.get(status, 0) + 1
        next_date = str(sheet.cell(row, 19).value or "")[:10]
        if next_date and next_date <= today_iso() and status != "已掌握":
            due += 1
    workbook.close()
    print(json.dumps({"total": total, "due": due, "statuses": counts}, ensure_ascii=False, indent=2))
    return 0


def parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="command", required=True)
    for command, help_text, function in (
        ("init", "Create the weakness workbook if it does not exist", cmd_init),
        ("flush", "Retry updates queued while Excel was locked", cmd_flush),
        ("summary", "Show weakness status counts", cmd_summary),
    ):
        item = sub.add_parser(command, help=help_text)
        item.add_argument("--workspace", required=True)
        item.set_defaults(func=function)
    upsert = sub.add_parser("upsert", help="Add or update weakness detail records")
    upsert.add_argument("--workspace", required=True)
    upsert.add_argument("--input", required=True)
    upsert.set_defaults(func=cmd_upsert)
    review = sub.add_parser("review", help="Append review history and update mastery")
    review.add_argument("--workspace", required=True)
    review.add_argument("--input", required=True)
    review.set_defaults(func=cmd_review)
    return ap


def main() -> int:
    args = parser().parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
