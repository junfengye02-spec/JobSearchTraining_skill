#!/usr/bin/env python3
"""Maintain verified job state and export the single job Excel workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as exc:  # pragma: no cover - environment dependent
    raise SystemExit("openpyxl is required. Install it with: python -m pip install openpyxl") from exc


RUNTIME_DIR = ".adaptive-interview-coach"
VOLATILE_QUERY_PREFIXES = (
    "utm_", "spm", "session", "token", "sid", "trace", "timestamp", "_t", "from",
)
REQUIRED_FIELDS = ("company", "title", "responsibilities", "source_url")
ALLOWED_DEACTIVATION_REASONS = {
    "deadline_passed": "截止日期已过",
    "removed": "职位已下线",
    "closed": "招聘已停止",
    "http_404": "来源页面返回404",
    "http_410": "来源页面返回410",
    "not_found": "官方来源明确不存在",
}

HEADERS = [
    "岗位名称", "岗位职责", "公司", "企业标签", "方向标签", "地点", "岗位性质",
    "地点优先级", "匹配度", "匹配依据", "能力缺口", "岗位亮点", "核实状态",
    "来源平台", "岗位链接", "首次发现", "最近确认", "截止日期", "届别",
    "招聘季节", "状态", "失效原因",
]
PROFILE_HEADERS = [
    "职责类别", "归纳职责", "出现岗位数", "覆盖比例", "优先级", "趋势", "关键词",
    "代表公司", "代表岗位", "简历已有证据", "简历能力缺口", "新增岗位影响", "画像更新时间",
]

NAVY = "18324A"
TEAL = "0F766E"
WHITE = "FFFFFF"
LIGHT_TEAL = "D9F0EC"
LIGHT_BLUE = "E7F0F7"
LIGHT_GRAY = "F3F4F6"
ORANGE = "FDE7C2"
RED = "FADBD8"
THIN = Side(style="thin", color="D1D5DB")


def now_iso() -> str:
    return datetime.now().astimezone().isoformat()


def today_iso() -> str:
    return date.today().isoformat()


def normalize_text(value: object) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"\s+", "", text)
    return text.translate(str.maketrans("（）【】", "()[]"))


def canonical_url(url: str) -> str:
    if not url:
        return ""
    parts = urlsplit(url.strip())
    kept = [
        (key, value)
        for key, value in parse_qsl(parts.query, keep_blank_values=True)
        if not any(key.lower().startswith(prefix) for prefix in VOLATILE_QUERY_PREFIXES)
    ]
    kept.sort()
    return urlunsplit((parts.scheme, parts.netloc, parts.path.rstrip("/"), urlencode(kept), ""))


def make_job_id(company: str, title: str, url: str) -> str:
    raw = f"{normalize_text(company)}|{normalize_text(title)}|{canonical_url(url)}"
    return "JOB-" + hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12].upper()


def fuzzy_key(company: str, title: str) -> str:
    return f"{normalize_text(company)}|{normalize_text(title)}"


def read_json(path: Path, default: object) -> object:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        return default


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    temp.replace(path)


def load_config(workspace: Path) -> dict:
    path = workspace / RUNTIME_DIR / "config.json"
    if not path.exists():
        raise SystemExit(f"Workspace is not initialized: {path}")
    return json.loads(path.read_text(encoding="utf-8"))


def resolve_config_path(workspace: Path, value: str) -> Path:
    path = Path(value)
    return path if path.is_absolute() else workspace / path


def state_path(workspace: Path, config: dict | None = None) -> Path:
    config = config or load_config(workspace)
    return resolve_config_path(workspace, config["files"]["job_state"])


def load_state(workspace: Path, config: dict | None = None) -> dict:
    config = config or load_config(workspace)
    return read_json(
        state_path(workspace, config),
        {"schema_version": 1, "updated_at": None, "jobs": {}},
    )


def save_state(workspace: Path, state: dict, config: dict | None = None) -> None:
    config = config or load_config(workspace)
    state["updated_at"] = now_iso()
    write_json(state_path(workspace, config), state)


def as_text(value: object) -> str:
    if isinstance(value, list):
        return "、".join(str(item).strip() for item in value if str(item).strip())
    return str(value or "").strip()


def verification_rank(value: str) -> int:
    text = normalize_text(value)
    if any(word in text for word in ("官方", "招聘官网", "official")):
        return 3
    if any(word in text for word in ("职位页", "招聘平台", "平台核实")):
        return 2
    if text:
        return 1
    return 0


def passes_filters(candidate: dict, config: dict) -> tuple[bool, str]:
    profile = config.get("profile", {})
    filters = config.get("filters", {})
    text = " ".join(
        as_text(candidate.get(field))
        for field in ("title", "responsibilities", "direction_tags", "highlight", "job_type")
    ).lower()
    positive = [item.lower() for item in filters.get("positive_keywords", []) if item]
    fuzzy = [item.lower() for item in filters.get("fuzzy_keywords", []) if item]
    negative = [item.lower() for item in filters.get("negative_keywords", []) if item]
    positive_hit = not positive or any(item in text for item in positive)
    fuzzy_hit = any(item in text for item in fuzzy)
    if positive and not positive_hit and not fuzzy_hit:
        return False, "未命中目标岗位方向"
    if not positive_hit and any(item in text for item in negative):
        return False, "命中排除方向"
    if not profile.get("include_internships", False):
        if any(word in text for word in ("实习", "intern")):
            return False, "配置未包含实习岗位"
    target = as_text(profile.get("target_grad_year"))
    target_year_match = re.search(r"(20\d{2})", target)
    years = [int(value) for value in re.findall(r"(20\d{2})届", text)]
    if target_year_match and years and max(years) < int(target_year_match.group(1)):
        return False, "岗位届别早于目标届别"
    return True, ""


def candidate_list(path: Path) -> list[dict]:
    value = read_json(path, [])
    if isinstance(value, dict):
        value = value.get("jobs", [])
    if not isinstance(value, list):
        raise SystemExit("Candidate input must be a JSON array or an object with a jobs array.")
    return [item for item in value if isinstance(item, dict)]


def prefer_new_value(existing: dict, candidate: dict, field: str) -> None:
    value = candidate.get(field)
    if value in (None, "", []):
        return
    current = existing.get(field)
    if current in (None, "", []) or len(as_text(value)) >= len(as_text(current)):
        existing[field] = value


def merge_candidates(workspace: Path, candidates: list[dict]) -> tuple[list[dict], list[dict], dict]:
    config = load_config(workspace)
    state = load_state(workspace, config)
    jobs = state.setdefault("jobs", {})
    fuzzy_index = {
        fuzzy_key(record.get("company", ""), record.get("title", "")): job_id
        for job_id, record in jobs.items()
    }
    new_records: list[dict] = []
    rejected: list[dict] = []
    today = today_iso()

    for candidate in candidates:
        missing = [field for field in REQUIRED_FIELDS if not as_text(candidate.get(field))]
        if missing:
            rejected.append({"candidate": candidate, "reason": "缺少必填字段: " + ", ".join(missing)})
            continue
        accepted, reason = passes_filters(candidate, config)
        if not accepted:
            rejected.append({"candidate": candidate, "reason": reason})
            continue

        company = as_text(candidate["company"])
        title = as_text(candidate["title"])
        source_url = as_text(candidate["source_url"])
        exact_id = make_job_id(company, title, source_url)
        existing_id = exact_id if exact_id in jobs else fuzzy_index.get(fuzzy_key(company, title))

        if existing_id:
            record = jobs[existing_id]
            for field in (
                "responsibilities", "city", "job_type", "enterprise_tag", "direction_tags",
                "match_level", "match_reason", "skill_gaps", "highlight", "source_platform",
                "deadline", "grad_year", "season",
            ):
                prefer_new_value(record, candidate, field)
            new_verification = as_text(candidate.get("verification"))
            if verification_rank(new_verification) >= verification_rank(as_text(record.get("verification"))):
                record["verification"] = new_verification
                record["source_url"] = source_url
                record["source_platform"] = as_text(candidate.get("source_platform"))
            record["canonical_url"] = canonical_url(record.get("source_url", source_url))
            record["last_confirmed"] = today
            record["active"] = True
            record["deactivated_at"] = None
            record["deactivation_reason"] = ""
            record["deactivation_evidence"] = ""
            continue

        record = {
            "job_id": exact_id,
            "company": company,
            "title": title,
            "responsibilities": as_text(candidate.get("responsibilities")),
            "city": as_text(candidate.get("city")) or "未注明",
            "job_type": as_text(candidate.get("job_type")),
            "enterprise_tag": as_text(candidate.get("enterprise_tag")),
            "direction_tags": candidate.get("direction_tags", []),
            "match_level": as_text(candidate.get("match_level")) or "待评估",
            "match_reason": as_text(candidate.get("match_reason")),
            "skill_gaps": as_text(candidate.get("skill_gaps")),
            "highlight": as_text(candidate.get("highlight")),
            "verification": as_text(candidate.get("verification")) or "待核实",
            "source_platform": as_text(candidate.get("source_platform")),
            "source_url": source_url,
            "canonical_url": canonical_url(source_url),
            "deadline": as_text(candidate.get("deadline")),
            "grad_year": as_text(candidate.get("grad_year")) or as_text(config.get("profile", {}).get("target_grad_year")),
            "season": as_text(candidate.get("season")) or as_text(config.get("profile", {}).get("target_season_label")),
            "first_seen": as_text(candidate.get("discovered_at")) or today,
            "last_confirmed": today,
            "active": True,
            "deactivated_at": None,
            "deactivation_reason": "",
            "deactivation_evidence": "",
        }
        jobs[exact_id] = record
        fuzzy_index[fuzzy_key(company, title)] = exact_id
        new_records.append(record.copy())

    save_state(workspace, state, config)
    return new_records, rejected, state


def cmd_merge(args: argparse.Namespace) -> int:
    workspace = Path(args.workspace).expanduser().resolve()
    candidates = candidate_list(Path(args.input).expanduser().resolve())
    new_records, rejected, state = merge_candidates(workspace, candidates)
    if args.new_output:
        write_json(Path(args.new_output).expanduser().resolve(), new_records)
    if args.rejected_output:
        write_json(Path(args.rejected_output).expanduser().resolve(), rejected)
    print(json.dumps({
        "input": len(candidates),
        "new": len(new_records),
        "rejected": len(rejected),
        "active_total": sum(1 for item in state["jobs"].values() if item.get("active", True)),
    }, ensure_ascii=False, indent=2))
    return 0


def find_record(jobs: dict, item: dict) -> tuple[str, dict] | tuple[None, None]:
    supplied_id = as_text(item.get("job_id"))
    if supplied_id and supplied_id in jobs:
        return supplied_id, jobs[supplied_id]
    company = as_text(item.get("company"))
    title = as_text(item.get("title"))
    url = as_text(item.get("source_url"))
    exact = make_job_id(company, title, url) if company and title else ""
    if exact and exact in jobs:
        return exact, jobs[exact]
    key = fuzzy_key(company, title)
    for job_id, record in jobs.items():
        if fuzzy_key(record.get("company", ""), record.get("title", "")) == key:
            return job_id, record
    return None, None


def cmd_deactivate(args: argparse.Namespace) -> int:
    workspace = Path(args.workspace).expanduser().resolve()
    config = load_config(workspace)
    state = load_state(workspace, config)
    requests = candidate_list(Path(args.input).expanduser().resolve())
    changed = []
    rejected = []
    for item in requests:
        reason = as_text(item.get("reason"))
        evidence = as_text(item.get("evidence"))
        if reason not in ALLOWED_DEACTIVATION_REASONS or not evidence:
            rejected.append({"request": item, "reason": "失效原因不受支持或缺少明确证据"})
            continue
        job_id, record = find_record(state.get("jobs", {}), item)
        if not record:
            rejected.append({"request": item, "reason": "未找到对应岗位"})
            continue
        record["active"] = False
        record["deactivated_at"] = today_iso()
        record["deactivation_reason"] = ALLOWED_DEACTIVATION_REASONS[reason]
        record["deactivation_evidence"] = evidence
        changed.append(job_id)
    save_state(workspace, state, config)
    print(json.dumps({"deactivated": changed, "rejected": rejected}, ensure_ascii=False, indent=2))
    return 0


def list_text(value: object) -> list[str]:
    if isinstance(value, list):
        return [as_text(item) for item in value if as_text(item)]
    text = as_text(value)
    return [text] if text else []


def profile_key(cluster: dict) -> str:
    return f"{normalize_text(cluster.get('category'))}|{normalize_text(cluster.get('summary'))}"


def build_duty_profile(state: dict, value: object) -> dict:
    if not isinstance(value, dict):
        raise SystemExit("Duty profile input must be a JSON object.")
    clusters = value.get("clusters", [])
    if not isinstance(clusters, list):
        raise SystemExit("Duty profile input must contain a clusters array.")

    active_jobs = {
        job_id: record
        for job_id, record in state.get("jobs", {}).items()
        if record.get("active", True) and as_text(record.get("responsibilities"))
    }
    active_count = len(active_jobs)
    previous = state.get("duty_profile", {})
    previous_by_key = {
        profile_key(cluster): cluster
        for cluster in previous.get("clusters", [])
        if isinstance(cluster, dict)
    }
    requested_new_ids = list_text(value.get("new_job_ids"))
    new_job_ids = [job_id for job_id in requested_new_ids if job_id in active_jobs]
    cleaned = []

    for raw in clusters:
        if not isinstance(raw, dict):
            continue
        category = as_text(raw.get("category"))
        summary = as_text(raw.get("summary"))
        if not category or not summary:
            continue
        try:
            frequency = int(raw.get("frequency", 0))
        except (TypeError, ValueError):
            frequency = 0
        frequency = max(0, min(frequency, active_count))
        representative_ids = [job_id for job_id in list_text(raw.get("representative_job_ids")) if job_id in active_jobs]
        representative_jobs = list_text(raw.get("representative_jobs"))
        if representative_ids:
            representative_jobs = [
                f"{active_jobs[job_id].get('company', '')}｜{active_jobs[job_id].get('title', '')}"
                for job_id in representative_ids
            ]
        companies = list_text(raw.get("companies"))
        if not companies and representative_ids:
            companies = list(dict.fromkeys(as_text(active_jobs[job_id].get("company")) for job_id in representative_ids))
        cluster = {
            "category": category,
            "summary": summary,
            "frequency": frequency,
            "coverage": round(frequency / active_count, 4) if active_count else 0,
            "priority": as_text(raw.get("priority")) or "中",
            "keywords": list_text(raw.get("keywords")),
            "companies": companies,
            "representative_job_ids": representative_ids,
            "representative_jobs": representative_jobs,
            "resume_evidence": as_text(raw.get("resume_evidence")),
            "resume_gap": as_text(raw.get("resume_gap")),
            "new_job_influence": as_text(raw.get("new_job_influence")),
        }
        old = previous_by_key.get(profile_key(cluster))
        if not old:
            trend = "新增"
        else:
            old_frequency = int(old.get("frequency", 0) or 0)
            trend = "上升" if frequency > old_frequency else "下降" if frequency < old_frequency else "稳定"
        cluster["trend"] = trend
        cleaned.append(cluster)

    cleaned.sort(key=lambda item: (-item["frequency"], normalize_text(item["category"]), normalize_text(item["summary"])))
    return {
        "generated_at": now_iso(),
        "source_job_count": active_count,
        "new_job_ids": new_job_ids,
        "clusters": cleaned,
    }


def cmd_profile(args: argparse.Namespace) -> int:
    workspace = Path(args.workspace).expanduser().resolve()
    config = load_config(workspace)
    state = load_state(workspace, config)
    value = read_json(Path(args.input).expanduser().resolve(), {})
    state["duty_profile"] = build_duty_profile(state, value)
    save_state(workspace, state, config)
    profile = state["duty_profile"]
    print(json.dumps({
        "source_job_count": profile["source_job_count"],
        "new_job_count": len(profile["new_job_ids"]),
        "clusters": len(profile["clusters"]),
        "generated_at": profile["generated_at"],
    }, ensure_ascii=False, indent=2))
    return 0


def location_rank(city: str, config: dict) -> int:
    preferences = config.get("profile", {}).get("location_preferences", [])
    for index, location in enumerate(preferences):
        if location and location in city:
            return index
    return len(preferences) + 1


def location_label(city: str, config: dict) -> str:
    preferences = config.get("profile", {}).get("location_preferences", [])
    hits = [location for location in preferences if location and location in city]
    return "、".join(hits) + "优先" if hits else "其他地区"


def match_rank(value: str) -> int:
    ranks = {"高度匹配": 0, "较高匹配": 1, "相近方向": 2, "待评估": 3, "待确认": 4}
    return ranks.get(value, 5)


def sorted_records(records: list[dict], config: dict) -> list[dict]:
    return sorted(
        records,
        key=lambda item: (
            location_rank(as_text(item.get("city")), config),
            match_rank(as_text(item.get("match_level"))),
            as_text(item.get("company")),
            as_text(item.get("title")),
        ),
    )


def add_job_sheet(workbook: Workbook, name: str, records: list[dict], config: dict, table_name: str) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(HEADERS)
    for record in sorted_records(records, config):
        city = as_text(record.get("city")) or "未注明"
        sheet.append([
            as_text(record.get("title")),
            as_text(record.get("responsibilities")),
            as_text(record.get("company")),
            as_text(record.get("enterprise_tag")),
            as_text(record.get("direction_tags")),
            city,
            as_text(record.get("job_type")),
            location_label(city, config),
            as_text(record.get("match_level")),
            as_text(record.get("match_reason")),
            as_text(record.get("skill_gaps")),
            as_text(record.get("highlight")),
            as_text(record.get("verification")),
            as_text(record.get("source_platform")),
            as_text(record.get("source_url")),
            as_text(record.get("first_seen")),
            as_text(record.get("last_confirmed")),
            as_text(record.get("deadline")),
            as_text(record.get("grad_year")),
            as_text(record.get("season")),
            "有效" if record.get("active", True) else "已失效",
            as_text(record.get("deactivation_reason")),
        ])
        link = sheet.cell(sheet.max_row, 15)
        if link.value:
            link.hyperlink = link.value
            link.style = "Hyperlink"

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=TEAL))
    sheet.row_dimensions[1].height = 32

    for row in sheet.iter_rows(min_row=2):
        rank = location_rank(as_text(row[5].value), config)
        color = LIGHT_TEAL if rank == 0 else LIGHT_BLUE if rank == 1 else WHITE
        if name == "已失效岗位":
            color = LIGHT_GRAY
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=color)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        if row[8].value in ("相近方向", "待评估"):
            row[8].fill = PatternFill("solid", fgColor=ORANGE)
        elif row[8].value == "待确认":
            row[8].fill = PatternFill("solid", fgColor=RED)
        row[0].font = Font(bold=True, color=NAVY)
        sheet.row_dimensions[row[0].row].height = 88

    widths = [38, 70, 22, 22, 28, 18, 18, 18, 14, 40, 38, 40, 18, 24, 48, 14, 14, 14, 12, 20, 12, 28]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    if records:
        table = Table(displayName=table_name, ref=f"A1:V{len(records) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        sheet.add_table(table)


def add_profile_sheet(workbook: Workbook, duty_profile: dict) -> None:
    sheet = workbook.create_sheet("职责画像")
    sheet.append(PROFILE_HEADERS)
    generated_at = as_text(duty_profile.get("generated_at"))
    for cluster in duty_profile.get("clusters", []):
        sheet.append([
            as_text(cluster.get("category")),
            as_text(cluster.get("summary")),
            int(cluster.get("frequency", 0) or 0),
            float(cluster.get("coverage", 0) or 0),
            as_text(cluster.get("priority")),
            as_text(cluster.get("trend")),
            as_text(cluster.get("keywords")),
            as_text(cluster.get("companies")),
            as_text(cluster.get("representative_jobs")),
            as_text(cluster.get("resume_evidence")),
            as_text(cluster.get("resume_gap")),
            as_text(cluster.get("new_job_influence")),
            generated_at,
        ])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=TEAL))
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        row[3].number_format = "0.0%"
        row[0].font = Font(bold=True, color=NAVY)
        if row[5].value == "新增" or row[5].value == "上升":
            row[5].fill = PatternFill("solid", fgColor=ORANGE)
        sheet.row_dimensions[row[0].row].height = 76
    widths = [20, 60, 14, 14, 12, 12, 28, 28, 38, 42, 42, 42, 24]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    if duty_profile.get("clusters"):
        table = Table(displayName="DutyProfile", ref=f"A1:M{len(duty_profile['clusters']) + 1}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=False)
        sheet.add_table(table)


def add_overview(workbook: Workbook, config: dict, active: list[dict], inactive: list[dict], duty_profile: dict) -> None:
    sheet = workbook.create_sheet("概览")
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 100
    sheet.append(["项目", "内容"])
    profile = config.get("profile", {})
    rows = [
        ("更新时间", now_iso()),
        ("目标届别", as_text(profile.get("target_grad_year"))),
        ("目标方向", as_text(profile.get("roles"))),
        ("招聘季节", as_text(profile.get("target_season_label"))),
        ("包含实习", "是" if profile.get("include_internships") else "否"),
        ("有效岗位数", len(active)),
        ("已失效岗位数", len(inactive)),
        ("职责画像条目", len(duty_profile.get("clusters", []))),
        ("职责画像更新时间", as_text(duty_profile.get("generated_at"))),
        ("使用说明", "面试训练优先读取每天重算的“职责画像”，再结合最新简历出题。投递前仍需打开岗位来源链接确认。"),
    ]
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row in sheet.iter_rows(min_row=2):
        row[0].font = Font(bold=True, color=NAVY)
        row[0].fill = PatternFill("solid", fgColor=LIGHT_GRAY)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row:
            cell.border = Border(bottom=THIN)
        sheet.row_dimensions[row[0].row].height = 38
    sheet.freeze_panes = "A2"


def save_workbook(workbook: Workbook, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    temp = output.with_name(output.name + ".tmp.xlsx")
    try:
        workbook.save(temp)
        os.replace(temp, output)
    except PermissionError as exc:
        temp.unlink(missing_ok=True)
        raise SystemExit(f"Excel workbook is open or locked; no copy was created: {output}") from exc
    finally:
        temp.unlink(missing_ok=True)


def export_workbook(workspace: Path) -> Path:
    config = load_config(workspace)
    state = load_state(workspace, config)
    records = list(state.get("jobs", {}).values())
    active = [item for item in records if item.get("active", True)]
    inactive = [item for item in records if not item.get("active", True)]
    duty_profile = state.get("duty_profile", {"clusters": []})
    workbook = Workbook()
    workbook.remove(workbook.active)
    add_job_sheet(workbook, "岗位总表", active, config, "ActiveJobs")
    add_profile_sheet(workbook, duty_profile)
    add_job_sheet(workbook, "已失效岗位", inactive, config, "InactiveJobs")
    add_overview(workbook, config, active, inactive, duty_profile)
    output = resolve_config_path(workspace, config["files"]["job_workbook"])
    save_workbook(workbook, output)
    load_workbook(output, read_only=True).close()
    return output


def ensure_workbook(workspace: Path) -> Path:
    config = load_config(workspace)
    output = resolve_config_path(workspace, config["files"]["job_workbook"])
    if not output.exists():
        return export_workbook(workspace)
    return output


def cmd_export(args: argparse.Namespace) -> int:
    output = export_workbook(Path(args.workspace).expanduser().resolve())
    print(json.dumps({"output": str(output), **state_summary(Path(args.workspace).expanduser().resolve())}, ensure_ascii=False, indent=2))
    return 0


def state_summary(workspace: Path) -> dict:
    config = load_config(workspace)
    state = load_state(workspace, config)
    values = list(state.get("jobs", {}).values())
    duty_profile = state.get("duty_profile", {})
    return {
        "active": sum(1 for item in values if item.get("active", True)),
        "inactive": sum(1 for item in values if not item.get("active", True)),
        "duty_profile_clusters": len(duty_profile.get("clusters", [])),
        "duty_profile_updated_at": duty_profile.get("generated_at"),
        "updated_at": state.get("updated_at"),
    }


def cmd_summary(args: argparse.Namespace) -> int:
    print(json.dumps(state_summary(Path(args.workspace).expanduser().resolve()), ensure_ascii=False, indent=2))
    return 0


def parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="command", required=True)

    merge = sub.add_parser("merge", help="Validate, filter, deduplicate, and merge discovered jobs")
    merge.add_argument("--workspace", required=True)
    merge.add_argument("--input", required=True)
    merge.add_argument("--new-output")
    merge.add_argument("--rejected-output")
    merge.set_defaults(func=cmd_merge)

    deactivate = sub.add_parser("deactivate", help="Deactivate jobs only with explicit supported evidence")
    deactivate.add_argument("--workspace", required=True)
    deactivate.add_argument("--input", required=True)
    deactivate.set_defaults(func=cmd_deactivate)

    profile = sub.add_parser("profile", help="Replace the aggregated duty profile after each job refresh")
    profile.add_argument("--workspace", required=True)
    profile.add_argument("--input", required=True)
    profile.set_defaults(func=cmd_profile)

    export = sub.add_parser("export", help="Rebuild the single formatted job workbook")
    export.add_argument("--workspace", required=True)
    export.set_defaults(func=cmd_export)

    summary = sub.add_parser("summary", help="Print active and inactive job counts")
    summary.add_argument("--workspace", required=True)
    summary.set_defaults(func=cmd_summary)
    return ap


def main() -> int:
    args = parser().parse_args()
    return args.func(args)


if __name__ == "__main__":
    raise SystemExit(main())
